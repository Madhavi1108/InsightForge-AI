"""Phase 33 (spec Phases 64-65, FR-24) - automated Excel & PDF reports.

DataFrame assembly (dataclass-list -> DataFrame helpers, the Executive
Summary rollup) is pure and gets direct unit tests, as does writing a
:class:`~src.reporting.ReportData` built entirely from stub data to real
``.xlsx``/``.pdf`` files on disk (no database - `write_excel_report`/
`write_pdf_report` only format already-fetched data). Fetching each section
from a real ``fact_sales``/analytics stack and the orchestrator wiring
(Stage 9, advisory - never affects ``pipeline_runs.status``) need real data,
so those are ``INSIGHTFORGE_PG_INTEGRATION=1``-gated (development rule 1:
never fake functionality), the same convention as Phases 20/21/26.
"""
from __future__ import annotations

import os
import subprocess
import sys
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from src.change_detection import ChangeRecord
from src.forecasting import ForecastPoint
from src.impact_analysis import BusinessImpactResult
from src.recommendations import Recommendation
from src.reporting import (
    SHEET_NAMES,
    ReportData,
    generate_reports,
    write_excel_report,
    write_pdf_report,
)
from src.root_cause import RootCauseResult

PROJECT_ROOT = Path(__file__).resolve().parents[1]

pg_integration = pytest.mark.skipif(
    os.environ.get("INSIGHTFORGE_PG_INTEGRATION") != "1",
    reason="set INSIGHTFORGE_PG_INTEGRATION=1 with a running PostgreSQL to enable",
)
SENTINEL = "__phase33_test__"


def test_phase33_files_exist():
    for p in (
        PROJECT_ROOT / "src" / "reporting.py",
        PROJECT_ROOT / "scripts" / "generate_report.py",
        PROJECT_ROOT / "docs" / "reporting.md",
    ):
        assert p.is_file(), f"missing {p}"


def test_import_has_no_side_effects():
    r = subprocess.run(
        [sys.executable, "-c", "import src.reporting"],
        cwd=PROJECT_ROOT, capture_output=True, text=True, check=False,
    )
    assert r.returncode == 0, r.stderr


def test_sheet_names_match_spec():
    assert SHEET_NAMES == (
        "Executive Summary", "KPIs", "Regions", "Categories", "Products",
        "Customers", "Anomalies", "Root Causes", "Impact", "Recommendations",
        "Forecast", "Data Quality",
    )


# --------------------------------------------------------------------------- #
# Fixture: a fully-stubbed ReportData - no database
# --------------------------------------------------------------------------- #
def _stub_report_data(*, empty: bool = False) -> ReportData:
    if empty:
        return ReportData(
            run_id=1, report_date="2026-09-08",
            kpis=pd.DataFrame(), regions=pd.DataFrame(), categories=pd.DataFrame(),
            products=pd.DataFrame(), customers=pd.DataFrame(), anomalies=pd.DataFrame(),
            data_quality=pd.DataFrame(), dq_report=None,
            errors={"KPIs": "KPIs unavailable: OperationalError"},
        )

    kpis = pd.DataFrame([
        {"order_date": "2026-09-07", "revenue": 10000.0, "profit": 2000.0,
         "margin_pct": 20.0, "orders": 50},
        {"order_date": "2026-09-08", "revenue": 12000.0, "profit": 2400.0,
         "margin_pct": 20.0, "orders": 55},
    ])
    anomalies = pd.DataFrame([
        {"metric": "revenue", "anomaly_date": "2026-09-08", "direction": "up",
         "severity": "HIGH", "confidence": 0.8},
    ])
    changes = [ChangeRecord(
        grain="day", metric="revenue", period_key="2026-09-08",
        previous_period_key="2026-09-07", current=12000.0, previous=10000.0,
        pct_change=20.0, direction="up", magnitude=20.0, significant=True,
    )]
    root_causes = [RootCauseResult(
        metric="revenue", change=2000.0, primary_driver="West > Electronics",
        contribution=1200.0, confidence=0.75, evidence={},
    )]
    impact = BusinessImpactResult(
        date="2026-09-08", expected_revenue=11000.0, actual_revenue=12000.0,
        revenue_gap=1000.0, expected_profit=2200.0, actual_profit=2400.0,
        profit_gap=200.0, revenue_at_risk=0.0, profit_at_risk=0.0,
        customers_affected=40, orders_affected=55,
    )
    recommendations = [Recommendation(
        rule_id="demand_surge", title="Scale fulfillment & inventory for demand growth",
        rationale="Revenue and orders both moved up.", metric="revenue",
        period_key="2026-09-08", direction="up", severity="HIGH", confidence=0.8,
        impact_value=1000.0, priority_score=60.0, priority_band="HIGH",
        linked_anomaly_id=None, detail={},
    )]
    forecasts = [ForecastPoint(
        metric="revenue", horizon_days=7, forecast_date="2026-09-09",
        forecast_value=12500.0, lower_bound=11000.0, upper_bound=14000.0,
        model="exponential_smoothing", mae=100.0, rmse=120.0, mape=1.5,
    )]

    return ReportData(
        run_id=42, report_date="2026-09-08",
        kpis=kpis, regions=pd.DataFrame([{"region": "West", "revenue": 12000.0}]),
        categories=pd.DataFrame([{"category": "Electronics", "revenue": 12000.0}]),
        products=pd.DataFrame([{"product_id": "P1", "revenue": 12000.0}]),
        customers=pd.DataFrame([{"customer_id": "C1", "revenue": 12000.0}]),
        anomalies=anomalies,
        data_quality=pd.DataFrame([{"dimension": "Overall", "score": 98.0, "passed": True}]),
        dq_report=None, changes=changes, root_causes=root_causes, impact=impact,
        recommendations=recommendations, forecasts=forecasts, errors={},
    )


# --------------------------------------------------------------------------- #
# Excel - pure formatting over stub data
# --------------------------------------------------------------------------- #
def test_write_excel_report_has_all_12_sheets(tmp_path):
    data = _stub_report_data()
    path = tmp_path / "report.xlsx"
    write_excel_report(data, path)
    assert path.is_file()
    wb = openpyxl.load_workbook(path)
    assert tuple(wb.sheetnames) == SHEET_NAMES


def test_write_excel_report_populates_recommendations_sheet(tmp_path):
    data = _stub_report_data()
    path = tmp_path / "report.xlsx"
    write_excel_report(data, path)
    wb = openpyxl.load_workbook(path)
    ws = wb["Recommendations"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert "title" in header
    second_row = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))]
    assert "Scale fulfillment" in second_row[header.index("title")]


def test_write_excel_report_empty_sections_render_note_not_missing_sheet(tmp_path):
    data = _stub_report_data(empty=True)
    path = tmp_path / "report.xlsx"
    write_excel_report(data, path)
    wb = openpyxl.load_workbook(path)
    assert tuple(wb.sheetnames) == SHEET_NAMES
    ws = wb["KPIs"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert header == ["note"]


# --------------------------------------------------------------------------- #
# PDF - pure formatting over stub data
# --------------------------------------------------------------------------- #
def test_write_pdf_report_produces_nonempty_file(tmp_path):
    data = _stub_report_data()
    path = tmp_path / "report.pdf"
    write_pdf_report(data, path)
    assert path.is_file()
    assert path.stat().st_size > 0
    assert path.read_bytes()[:4] == b"%PDF"


def test_write_pdf_report_handles_empty_sections_without_error(tmp_path):
    data = _stub_report_data(empty=True)
    path = tmp_path / "report.pdf"
    write_pdf_report(data, path)
    assert path.is_file()
    assert path.read_bytes()[:4] == b"%PDF"


# --------------------------------------------------------------------------- #
# integration - real PostgreSQL + orchestrator wiring
# --------------------------------------------------------------------------- #
def _apply_schema():
    subprocess.run(
        [sys.executable, str(PROJECT_ROOT / "scripts" / "apply_schema.py")],
        cwd=PROJECT_ROOT, capture_output=True, text=True, check=False,
    )


def _int_paths(tmp_path):
    from src.config import PipelinePaths
    fields = ("incoming", "raw", "processed", "rejected", "archive", "logs", "reports")
    p = PipelinePaths(**{f: tmp_path / f for f in fields})
    p.ensure()
    return p


@pg_integration
def test_orchestrator_writes_reports_without_affecting_status(tmp_path, monkeypatch):
    pytest.importorskip("psycopg2")
    _apply_schema()
    from scripts import generate_dataset as gd
    from src import orchestrator
    from src.database import Database

    p = _int_paths(tmp_path)
    monkeypatch.setattr(orchestrator, "get_paths", lambda: p)
    db = Database()

    df = gd.generate_dataset(rows=2000, seed=33)
    dates = sorted(df["Order_Date"].unique())[:15]
    codes = []
    for i, d in enumerate(dates):
        day_df = df[df["Order_Date"] == d].copy()
        name = f"{SENTINEL}_{i}.csv"
        day_df.to_csv(p.incoming / name, index=False, date_format="%Y-%m-%d")
        codes.append(orchestrator.run_file(p.incoming / name))

    run_ids = [r["run_id"] for r in db.fetch_all(
        "SELECT run_id FROM pipeline_runs WHERE file_name LIKE :pat ORDER BY run_id",
        {"pat": SENTINEL + "%"},
    )]
    assert all(c in (0, 9) for c in codes)

    xlsx_files = list(p.reports.glob("InsightForge_Report_*.xlsx"))
    pdf_files = list(p.reports.glob("InsightForge_Executive_Report_*.pdf"))
    assert xlsx_files, "orchestrator should have written at least one Excel report"
    assert pdf_files, "orchestrator should have written at least one PDF report"

    runs = db.fetch_all(
        "SELECT run_id, status, stage_metrics FROM pipeline_runs WHERE run_id = ANY(:ids)",
        {"ids": run_ids},
    )
    for row in runs:
        assert row["status"] in ("SUCCESS", "WARNING", "FAILED")
        assert "report" in row["stage_metrics"]


@pg_integration
def test_generate_reports_standalone_reads_persisted_dq(tmp_path, monkeypatch):
    pytest.importorskip("psycopg2")
    _apply_schema()
    from scripts import generate_dataset as gd
    from src import orchestrator
    from src.database import Database

    p = _int_paths(tmp_path)
    monkeypatch.setattr(orchestrator, "get_paths", lambda: p)
    db = Database()

    df = gd.generate_dataset(rows=1500, seed=34)
    d = sorted(df["Order_Date"].unique())[0]
    day_df = df[df["Order_Date"] == d].copy()
    name = f"{SENTINEL}_standalone.csv"
    day_df.to_csv(p.incoming / name, index=False, date_format="%Y-%m-%d")

    code = orchestrator.run_file(p.incoming / name)
    assert code in (0, 9)

    run_id = db.fetch_one(
        "SELECT run_id FROM pipeline_runs WHERE file_name = :n", {"n": name},
    )["run_id"]

    # Standalone call (no in-memory DqReport) - Data Quality sheet must come
    # back from the persisted data_quality_results rows the run just wrote.
    excel_path, pdf_path = generate_reports(db, run_id, reports_dir=p.reports)
    assert excel_path.is_file()
    assert pdf_path.is_file()

    wb = openpyxl.load_workbook(excel_path)
    ws = wb["Data Quality"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert header != ["note"], "Data Quality sheet should have real rows, not an empty note"
